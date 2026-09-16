
import os
import csv
import io
import calendar
import secrets
from datetime import datetime, date, timedelta
from decimal import Decimal
from functools import wraps

from flask import Flask, render_template, request, redirect, url_for, flash, abort, Response, send_file
from flask_sqlalchemy import SQLAlchemy
from flask_login import (
    LoginManager, UserMixin, login_user, logout_user,
    login_required, current_user
)
from flask_wtf.csrf import CSRFProtect
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy import text

BASE_DIR = os.path.abspath(os.path.dirname(__file__))

app = Flask(__name__)

database_url = os.getenv(
    "DATABASE_URL",
    "sqlite:///" + os.path.join(BASE_DIR, "boston.db")
)
# Neon normalmente fornece postgresql://.
# Forçamos o driver psycopg v3, que é o instalado em requirements.txt.
if database_url.startswith("postgres://"):
    database_url = database_url.replace("postgres://", "postgresql+psycopg://", 1)
elif database_url.startswith("postgresql://"):
    database_url = database_url.replace("postgresql://", "postgresql+psycopg://", 1)

app.config["SECRET_KEY"] = os.getenv("SECRET_KEY", "CHANGE-ME-IN-PRODUCTION")
app.config["SQLALCHEMY_DATABASE_URI"] = database_url
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

# Neon é um PostgreSQL serverless: conexões podem ser recicladas quando o banco
# entra/sai de repouso. Estas opções evitam reutilizar conexões quebradas.
if database_url.startswith("postgresql"):
    app.config["SQLALCHEMY_ENGINE_OPTIONS"] = {
        "pool_pre_ping": True,
        "pool_recycle": 180,
        "pool_use_lifo": True,
        "pool_size": 3,
        "max_overflow": 2,
    }

# Cookies mais seguros quando estiver publicado com HTTPS.
is_production = os.getenv("RENDER", "").lower() == "true" or os.getenv("APP_ENV", "").lower() == "production"
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = is_production

db = SQLAlchemy(app)
csrf = CSRFProtect(app)
login_manager = LoginManager(app)
login_manager.login_view = "login"
login_manager.login_message = "Entre na sua conta para continuar."
login_manager.login_message_category = "info"


class User(UserMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(120), nullable=False)
    email = db.Column(db.String(180), unique=True, nullable=False, index=True)
    phone = db.Column(db.String(30))
    password_hash = db.Column(db.String(255), nullable=False)
    role = db.Column(db.String(20), default="client", nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)

    appointments = db.relationship("Appointment", backref="client", lazy=True)
    vehicles = db.relationship("Vehicle", backref="owner", lazy=True, cascade="all, delete-orphan")

    def set_password(self, password):
        self.password_hash = generate_password_hash(password)

    def check_password(self, password):
        return check_password_hash(self.password_hash, password)


class Vehicle(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    brand = db.Column(db.String(80), nullable=False)
    model = db.Column(db.String(100), nullable=False)
    year = db.Column(db.String(10))
    plate = db.Column(db.String(20))
    color = db.Column(db.String(50))
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)

    @property
    def display_name(self):
        parts = [self.brand, self.model]
        if self.year:
            parts.append(self.year)
        return " ".join(parts)


class ClientPlan(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    service_id = db.Column(db.Integer, db.ForeignKey("service.id"), nullable=True)
    plan_name = db.Column(db.String(120), nullable=False, default="Plano recorrente")
    frequency = db.Column(db.String(30), nullable=False, default="mensal")
    amount = db.Column(db.Numeric(10, 2), nullable=False, default=0)
    next_date = db.Column(db.Date, nullable=True)
    active = db.Column(db.Boolean, default=True, nullable=False)
    notes = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)

    client = db.relationship("User", backref=db.backref("plans", lazy=True, cascade="all, delete-orphan"))
    service = db.relationship("Service")


class Service(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(140), nullable=False)
    description = db.Column(db.Text)
    price = db.Column(db.Numeric(10, 2), nullable=False, default=0)
    duration_minutes = db.Column(db.Integer, default=60)
    active = db.Column(db.Boolean, default=True, nullable=False)

    appointments = db.relationship("Appointment", backref="service", lazy=True)


class Employee(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(120), nullable=False)
    role_name = db.Column(db.String(100), default="Lavador")
    phone = db.Column(db.String(30))
    commission_percent = db.Column(db.Numeric(5, 2), default=0, nullable=False)
    active = db.Column(db.Boolean, default=True, nullable=False)
    notes = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)

    appointments = db.relationship("Appointment", backref="employee", lazy=True)


class Appointment(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    service_id = db.Column(db.Integer, db.ForeignKey("service.id"), nullable=False)
    employee_id = db.Column(db.Integer, db.ForeignKey("employee.id"), nullable=True)

    vehicle = db.Column(db.String(140), nullable=False)
    plate = db.Column(db.String(20))
    scheduled_date = db.Column(db.Date, nullable=False)
    scheduled_time = db.Column(db.String(10), nullable=False)
    status = db.Column(db.String(30), default="agendado", nullable=False)
    amount = db.Column(db.Numeric(10, 2), nullable=False, default=0)
    payment_method = db.Column(db.String(50), default="Não informado", nullable=False)
    completed_at = db.Column(db.DateTime, nullable=True)
    notes = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)


class Transaction(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    kind = db.Column(db.String(20), nullable=False)  # income / expense
    category = db.Column(db.String(80), nullable=False)
    description = db.Column(db.String(220), nullable=False)
    amount = db.Column(db.Numeric(12, 2), nullable=False)
    payment_method = db.Column(db.String(50), default="Não informado")
    date = db.Column(db.Date, default=date.today, nullable=False)
    appointment_id = db.Column(db.Integer, nullable=True, index=True)
    automatic = db.Column(db.Boolean, default=False, nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)


class CashDay(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    date = db.Column(db.Date, unique=True, nullable=False, index=True)
    opening_balance = db.Column(db.Numeric(12, 2), default=0, nullable=False)
    closing_balance = db.Column(db.Numeric(12, 2), nullable=True)
    notes = db.Column(db.Text)
    closed = db.Column(db.Boolean, default=False, nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)


@login_manager.user_loader
def load_user(user_id):
    return db.session.get(User, int(user_id))


def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not current_user.is_authenticated:
            return redirect(url_for("login"))
        if current_user.role != "admin":
            abort(403)
        return view(*args, **kwargs)
    return wrapped


@app.template_filter("brl")
def brl(value):
    try:
        v = Decimal(value or 0)
        s = f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        return f"R$ {s}"
    except Exception:
        return "R$ 0,00"


def decimal_from_form(value, default="0"):
    try:
        return Decimal(str(value or default).replace(",", "."))
    except Exception:
        return Decimal(default)


def month_bounds(month_value=None):
    if month_value:
        try:
            year, month = [int(x) for x in month_value.split("-")]
        except Exception:
            today = date.today()
            year, month = today.year, today.month
    else:
        today = date.today()
        year, month = today.year, today.month

    last = calendar.monthrange(year, month)[1]
    return date(year, month, 1), date(year, month, last), f"{year:04d}-{month:02d}"


def transaction_totals(start_date, end_date):
    income = db.session.query(db.func.coalesce(db.func.sum(Transaction.amount), 0)).filter(
        Transaction.kind == "income",
        Transaction.date >= start_date,
        Transaction.date <= end_date
    ).scalar()

    expenses = db.session.query(db.func.coalesce(db.func.sum(Transaction.amount), 0)).filter(
        Transaction.kind == "expense",
        Transaction.date >= start_date,
        Transaction.date <= end_date
    ).scalar()

    income = Decimal(income or 0)
    expenses = Decimal(expenses or 0)
    return income, expenses, income - expenses


def payment_breakdown(start_date, end_date, kind="income"):
    rows = db.session.query(
        Transaction.payment_method,
        db.func.coalesce(db.func.sum(Transaction.amount), 0)
    ).filter(
        Transaction.kind == kind,
        Transaction.date >= start_date,
        Transaction.date <= end_date
    ).group_by(Transaction.payment_method).all()

    return {
        (method or "Não informado"): Decimal(total or 0)
        for method, total in rows
    }


def physical_cash_totals(target_date):
    cash_income = db.session.query(db.func.coalesce(db.func.sum(Transaction.amount), 0)).filter(
        Transaction.kind == "income",
        Transaction.date == target_date,
        Transaction.payment_method == "Dinheiro"
    ).scalar()

    cash_expense = db.session.query(db.func.coalesce(db.func.sum(Transaction.amount), 0)).filter(
        Transaction.kind == "expense",
        Transaction.date == target_date,
        Transaction.payment_method == "Dinheiro"
    ).scalar()

    return Decimal(cash_income or 0), Decimal(cash_expense or 0)


def manual_payment_breakdown(target_date):
    rows = db.session.query(
        Transaction.payment_method,
        db.func.coalesce(db.func.sum(Transaction.amount), 0)
    ).filter(
        Transaction.kind == "income",
        Transaction.date == target_date,
        Transaction.category == "Recebimento manual"
    ).group_by(Transaction.payment_method).all()

    return {
        (method or "Não informado"): Decimal(total or 0)
        for method, total in rows
    }


def advance_recurring_date(current_date, frequency):
    base = current_date or date.today()

    if frequency == "semanal":
        return base + timedelta(days=7)
    if frequency == "quinzenal":
        return base + timedelta(days=14)
    if frequency == "mensal":
        year = base.year + (base.month // 12)
        month = 1 if base.month == 12 else base.month + 1
        if base.month != 12:
            year = base.year
        day = min(base.day, calendar.monthrange(year, month)[1])
        return date(year, month, day)
    if frequency == "bimestral":
        total = base.year * 12 + (base.month - 1) + 2
        year, month0 = divmod(total, 12)
        month = month0 + 1
        day = min(base.day, calendar.monthrange(year, month)[1])
        return date(year, month, day)

    return base


def sync_appointment_finance(appt):
    income_tx = Transaction.query.filter_by(
        appointment_id=appt.id, kind="income", automatic=True
    ).first()
    commission_tx = Transaction.query.filter_by(
        appointment_id=appt.id, kind="expense", category="Comissão", automatic=True
    ).first()

    if appt.status == "concluido":
        if not appt.completed_at:
            appt.completed_at = datetime.utcnow()

        completion_date = appt.completed_at.date()
        payment_method = appt.payment_method or "Não informado"

        if not income_tx:
            income_tx = Transaction(
                kind="income",
                category="Serviço",
                description=f"{appt.service.name} - {appt.client.name}",
                amount=appt.amount,
                payment_method=payment_method,
                date=completion_date,
                appointment_id=appt.id,
                automatic=True,
            )
            db.session.add(income_tx)
        else:
            income_tx.amount = appt.amount
            income_tx.description = f"{appt.service.name} - {appt.client.name}"
            income_tx.payment_method = payment_method
            income_tx.date = completion_date

        if appt.employee and Decimal(appt.employee.commission_percent or 0) > 0:
            commission = (
                Decimal(appt.amount or 0)
                * Decimal(appt.employee.commission_percent or 0)
                / Decimal("100")
            ).quantize(Decimal("0.01"))

            if not commission_tx:
                commission_tx = Transaction(
                    kind="expense",
                    category="Comissão",
                    description=f"Comissão - {appt.employee.name} - {appt.service.name}",
                    amount=commission,
                    payment_method="Interno",
                    date=completion_date,
                    appointment_id=appt.id,
                    automatic=True,
                )
                db.session.add(commission_tx)
            else:
                commission_tx.amount = commission
                commission_tx.description = f"Comissão - {appt.employee.name} - {appt.service.name}"
                commission_tx.date = completion_date
        elif commission_tx:
            db.session.delete(commission_tx)
    else:
        appt.completed_at = None
        if income_tx:
            db.session.delete(income_tx)
        if commission_tx:
            db.session.delete(commission_tx)


@app.route("/")
def home():
    services = Service.query.filter_by(active=True).order_by(Service.name.asc()).all()
    return render_template("index.html", services=services)


@app.route("/cadastro", methods=["GET", "POST"])
def register():
    if current_user.is_authenticated:
        return redirect(url_for("client_dashboard"))
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        email = request.form.get("email", "").strip().lower()
        phone = request.form.get("phone", "").strip()
        password = request.form.get("password", "")

        if not name or not email or len(password) < 6:
            flash("Preencha nome, e-mail e uma senha com pelo menos 6 caracteres.", "danger")
        elif User.query.filter_by(email=email).first():
            flash("Este e-mail já está cadastrado.", "danger")
        else:
            user = User(name=name, email=email, phone=phone, role="client", password_hash="")
            user.set_password(password)
            db.session.add(user)
            db.session.commit()
            login_user(user)
            flash("Conta criada com sucesso.", "success")
            return redirect(url_for("client_dashboard"))

    return render_template("register.html")


@app.route("/entrar", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated:
        return redirect(url_for("admin_dashboard" if current_user.role == "admin" else "client_dashboard"))

    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        user = User.query.filter_by(email=email).first()

        if user and user.check_password(password):
            login_user(user)
            flash("Login realizado.", "success")
            return redirect(url_for("admin_dashboard" if user.role == "admin" else "client_dashboard"))

        flash("E-mail ou senha inválidos.", "danger")

    return render_template("login.html")


@app.route("/sair")
@login_required
def logout():
    logout_user()
    return redirect(url_for("home"))


@app.route("/cliente")
@login_required
def client_dashboard():
    if current_user.role == "admin":
        return redirect(url_for("admin_dashboard"))

    upcoming = Appointment.query.filter(
        Appointment.user_id == current_user.id,
        Appointment.scheduled_date >= date.today()
    ).order_by(Appointment.scheduled_date.asc(), Appointment.scheduled_time.asc()).all()

    history = Appointment.query.filter_by(user_id=current_user.id).order_by(
        Appointment.created_at.desc()
    ).all()

    services = Service.query.filter_by(active=True).order_by(Service.name.asc()).all()
    vehicles = Vehicle.query.filter_by(user_id=current_user.id).order_by(Vehicle.created_at.desc()).all()

    return render_template(
        "client_dashboard.html",
        upcoming=upcoming,
        history=history,
        services=services,
        vehicles=vehicles,
    )


@app.route("/cliente/veiculos", methods=["POST"])
@login_required
def client_add_vehicle():
    if current_user.role != "client":
        abort(403)

    brand = request.form.get("brand", "").strip()
    model = request.form.get("model", "").strip()
    if not brand or not model:
        flash("Informe marca e modelo do veículo.", "danger")
        return redirect(url_for("client_dashboard"))

    vehicle = Vehicle(
        user_id=current_user.id,
        brand=brand,
        model=model,
        year=request.form.get("year", "").strip(),
        plate=request.form.get("plate", "").strip().upper(),
        color=request.form.get("color", "").strip(),
    )
    db.session.add(vehicle)
    db.session.commit()
    flash("Veículo salvo.", "success")
    return redirect(url_for("client_dashboard"))


@app.route("/cliente/veiculos/<int:vehicle_id>/excluir", methods=["POST"])
@login_required
def client_delete_vehicle(vehicle_id):
    vehicle = db.session.get(Vehicle, vehicle_id)
    if not vehicle or vehicle.user_id != current_user.id:
        abort(404)
    db.session.delete(vehicle)
    db.session.commit()
    flash("Veículo removido.", "success")
    return redirect(url_for("client_dashboard"))


@app.route("/cliente/agendar", methods=["POST"])
@login_required
def client_book():
    if current_user.role != "client":
        abort(403)

    service = db.session.get(Service, int(request.form.get("service_id", "0") or 0))
    vehicle_id = int(request.form.get("vehicle_id", "0") or 0)
    vehicle_obj = db.session.get(Vehicle, vehicle_id) if vehicle_id else None

    if vehicle_obj and vehicle_obj.user_id != current_user.id:
        vehicle_obj = None

    manual_vehicle = request.form.get("vehicle", "").strip()
    vehicle_name = vehicle_obj.display_name if vehicle_obj else manual_vehicle
    plate = (vehicle_obj.plate if vehicle_obj else request.form.get("plate", "")).strip().upper()
    notes = request.form.get("notes", "").strip()
    date_str = request.form.get("date", "")
    time_str = request.form.get("time", "")

    if not service or not vehicle_name or not date_str or not time_str:
        flash("Preencha todos os campos obrigatórios do agendamento.", "danger")
        return redirect(url_for("client_dashboard"))

    try:
        scheduled_date = datetime.strptime(date_str, "%Y-%m-%d").date()
    except ValueError:
        flash("Data inválida.", "danger")
        return redirect(url_for("client_dashboard"))

    if scheduled_date < date.today():
        flash("Escolha uma data futura.", "danger")
        return redirect(url_for("client_dashboard"))

    appt = Appointment(
        user_id=current_user.id,
        service_id=service.id,
        vehicle=vehicle_name,
        plate=plate,
        scheduled_date=scheduled_date,
        scheduled_time=time_str,
        amount=service.price,
        notes=notes,
        status="agendado",
    )
    db.session.add(appt)
    db.session.commit()
    flash("Agendamento criado. A administração poderá confirmar pelo painel.", "success")
    return redirect(url_for("client_dashboard"))


@app.route("/admin")
@admin_required
def admin_dashboard():
    today = date.today()
    month_start, month_end, _ = month_bounds()

    today_income, today_expenses, today_profit = transaction_totals(today, today)
    month_income, month_expenses, month_profit = transaction_totals(month_start, month_end)

    cash_day = CashDay.query.filter_by(date=today).first()
    opening_balance = Decimal(cash_day.opening_balance or 0) if cash_day else Decimal("0")
    cash_income, cash_expense = physical_cash_totals(today)
    physical_cash_balance = opening_balance + cash_income - cash_expense

    payment_methods = payment_breakdown(today, today, "income")
    manual_payments = manual_payment_breakdown(today)

    upcoming = Appointment.query.filter(Appointment.scheduled_date >= today).count()
    open_count = Appointment.query.filter(
        Appointment.status.in_(["agendado", "confirmado", "em_atendimento"])
    ).count()
    client_count = User.query.filter_by(role="client").count()

    today_appointments = Appointment.query.filter(
        Appointment.scheduled_date == today,
        Appointment.status != "cancelado"
    ).order_by(Appointment.scheduled_time.asc()).all()

    completed_today = Appointment.query.filter(
        Appointment.status == "concluido",
        Appointment.completed_at.isnot(None),
        db.func.date(Appointment.completed_at) == today.isoformat()
    ).count()

    employees = Employee.query.filter_by(active=True).order_by(Employee.name.asc()).all()

    recent_transactions = Transaction.query.order_by(
        Transaction.date.desc(), Transaction.id.desc()
    ).limit(8).all()

    return render_template(
        "admin_dashboard.html",
        today_income=today_income,
        today_expenses=today_expenses,
        today_profit=today_profit,
        month_income=month_income,
        month_expenses=month_expenses,
        month_profit=month_profit,
        opening_balance=opening_balance,
        cash_income=cash_income,
        cash_expense=cash_expense,
        physical_cash_balance=physical_cash_balance,
        payment_methods=payment_methods,
        manual_payments=manual_payments,
        upcoming=upcoming,
        open_count=open_count,
        client_count=client_count,
        completed_today=completed_today,
        today_appointments=today_appointments,
        employees=employees,
        recent_transactions=recent_transactions,
    )


@app.route("/admin/agendamentos")
@admin_required
def admin_appointments():
    status = request.args.get("status", "").strip()
    q = Appointment.query

    if status:
        q = q.filter_by(status=status)

    appointments = q.order_by(
        Appointment.scheduled_date.desc(),
        Appointment.scheduled_time.desc()
    ).all()

    employees = Employee.query.filter_by(active=True).order_by(Employee.name.asc()).all()

    return render_template(
        "admin_appointments.html",
        appointments=appointments,
        selected_status=status,
        employees=employees,
    )


@app.route("/admin/agendamentos/<int:appointment_id>/status", methods=["POST"])
@admin_required
def admin_appointment_status(appointment_id):
    appt = db.session.get(Appointment, appointment_id)
    if not appt:
        abort(404)

    new_status = request.form.get("status", "")
    allowed = {"agendado", "confirmado", "em_atendimento", "concluido", "cancelado"}
    if new_status not in allowed:
        flash("Status inválido.", "danger")
        return redirect(request.referrer or url_for("admin_appointments"))

    employee_id = int(request.form.get("employee_id", "0") or 0)
    appt.employee_id = employee_id or None

    amount_raw = request.form.get("amount")
    if amount_raw not in (None, ""):
        amount = decimal_from_form(amount_raw)
        if amount < 0:
            flash("O valor final não pode ser negativo.", "danger")
            return redirect(request.referrer or url_for("admin_appointments"))
        appt.amount = amount

    payment_method = request.form.get("payment_method", "").strip()
    if payment_method:
        appt.payment_method = payment_method

    appt.status = new_status
    if new_status == "concluido" and not appt.completed_at:
        appt.completed_at = datetime.utcnow()

    sync_appointment_finance(appt)
    db.session.commit()

    if new_status == "concluido":
        flash(
            f"Atendimento concluído: {brl(appt.amount)} entrou no movimento de hoje.",
            "success"
        )
    else:
        flash("Agendamento atualizado.", "success")

    return redirect(request.referrer or url_for("admin_appointments"))


@app.route("/admin/agendamentos/<int:appointment_id>/concluir", methods=["POST"])
@admin_required
def admin_quick_complete(appointment_id):
    appt = db.session.get(Appointment, appointment_id)
    if not appt:
        abort(404)

    amount = decimal_from_form(request.form.get("amount", appt.amount))
    if amount < 0:
        flash("O valor final não pode ser negativo.", "danger")
        return redirect(request.referrer or url_for("admin_dashboard"))

    employee_id = int(request.form.get("employee_id", "0") or 0)
    payment_method = request.form.get("payment_method", "Pix").strip() or "Pix"

    appt.amount = amount
    appt.employee_id = employee_id or appt.employee_id
    appt.payment_method = payment_method
    appt.status = "concluido"
    appt.completed_at = datetime.utcnow()

    sync_appointment_finance(appt)
    db.session.commit()

    flash(
        f"Concluído! {brl(appt.amount)} foi adicionado ao movimento diário via {payment_method}.",
        "success"
    )
    return redirect(request.referrer or url_for("admin_dashboard"))


@app.route("/admin/financeiro", methods=["GET", "POST"])
@admin_required
def admin_finance():
    if request.method == "POST":
        kind = request.form.get("kind", "")
        category = request.form.get("category", "").strip()
        description = request.form.get("description", "").strip()
        payment_method = request.form.get("payment_method", "").strip() or "Não informado"
        amount = decimal_from_form(request.form.get("amount", "0"))
        date_str = request.form.get("date", "")

        if kind not in {"income", "expense"} or not category or not description or amount <= 0:
            flash("Preencha os dados corretamente.", "danger")
        else:
            try:
                tx_date = datetime.strptime(date_str, "%Y-%m-%d").date() if date_str else date.today()
                db.session.add(Transaction(
                    kind=kind,
                    category=category,
                    description=description,
                    amount=amount,
                    payment_method=payment_method,
                    date=tx_date,
                    automatic=False,
                ))
                db.session.commit()
                flash("Lançamento salvo.", "success")
                return redirect(url_for("admin_finance"))
            except Exception:
                db.session.rollback()
                flash("Valor ou data inválidos.", "danger")

    month_value = request.args.get("month", "")
    start_date, end_date, selected_month = month_bounds(month_value)
    transactions = Transaction.query.filter(
        Transaction.date >= start_date,
        Transaction.date <= end_date
    ).order_by(Transaction.date.desc(), Transaction.id.desc()).all()

    total_income, total_expense, total_profit = transaction_totals(start_date, end_date)

    return render_template(
        "admin_finance.html",
        transactions=transactions,
        total_income=total_income,
        total_expense=total_expense,
        total_profit=total_profit,
        selected_month=selected_month,
    )


@app.route("/admin/financeiro/<int:transaction_id>/excluir", methods=["POST"])
@admin_required
def admin_delete_transaction(transaction_id):
    tx = db.session.get(Transaction, transaction_id)
    if not tx:
        abort(404)

    if tx.automatic:
        flash("Lançamentos automáticos devem ser alterados pelo atendimento que os gerou.", "danger")
        return redirect(url_for("admin_finance"))

    db.session.delete(tx)
    db.session.commit()
    flash("Lançamento excluído.", "success")
    return redirect(url_for("admin_finance"))


@app.route("/admin/caixa", methods=["GET", "POST"])
@admin_required
def admin_cash():
    selected = request.args.get("date") or date.today().isoformat()

    try:
        selected_date = datetime.strptime(selected, "%Y-%m-%d").date()
    except ValueError:
        selected_date = date.today()

    cash = CashDay.query.filter_by(date=selected_date).first()

    if request.method == "POST":
        opening = decimal_from_form(request.form.get("opening_balance", "0"))
        notes = request.form.get("notes", "").strip()

        if not cash:
            cash = CashDay(date=selected_date, opening_balance=opening, notes=notes)
            db.session.add(cash)
        else:
            cash.opening_balance = opening
            cash.notes = notes

        db.session.commit()
        flash("Caixa do dia salvo.", "success")
        return redirect(url_for("admin_cash", date=selected_date.isoformat()))

    income, expenses, operational_result = transaction_totals(selected_date, selected_date)
    payment_methods = payment_breakdown(selected_date, selected_date, "income")
    manual_payments = manual_payment_breakdown(selected_date)

    opening = Decimal(cash.opening_balance or 0) if cash else Decimal("0")
    cash_income, cash_expense = physical_cash_totals(selected_date)
    expected_cash = opening + cash_income - cash_expense

    difference = None
    if cash and cash.closing_balance is not None:
        difference = Decimal(cash.closing_balance) - expected_cash

    transactions = Transaction.query.filter_by(date=selected_date).order_by(
        Transaction.id.desc()
    ).all()

    completed_services = Appointment.query.filter(
        Appointment.status == "concluido",
        Appointment.completed_at.isnot(None),
        db.func.date(Appointment.completed_at) == selected_date.isoformat()
    ).order_by(Appointment.completed_at.desc()).all()

    return render_template(
        "admin_cash.html",
        cash=cash,
        selected_date=selected_date,
        income=income,
        expenses=expenses,
        operational_result=operational_result,
        opening=opening,
        cash_income=cash_income,
        cash_expense=cash_expense,
        expected=expected_cash,
        difference=difference,
        payment_methods=payment_methods,
        manual_payments=manual_payments,
        transactions=transactions,
        completed_services=completed_services,
    )


@app.route("/admin/caixa/recebimento", methods=["POST"])
@admin_required
def admin_set_manual_payment():
    try:
        target_date = datetime.strptime(
            request.form.get("date", date.today().isoformat()),
            "%Y-%m-%d"
        ).date()
    except ValueError:
        target_date = date.today()

    method = request.form.get("payment_method", "").strip()
    allowed_methods = {"Pix", "Dinheiro", "Cartão", "Transferência", "Outro"}
    if method not in allowed_methods:
        flash("Forma de pagamento inválida.", "danger")
        return redirect(request.referrer or url_for("admin_cash"))

    amount = decimal_from_form(request.form.get("amount", "0"))
    if amount < 0:
        flash("O valor não pode ser negativo.", "danger")
        return redirect(request.referrer or url_for("admin_cash", date=target_date.isoformat()))

    existing = Transaction.query.filter_by(
        kind="income",
        category="Recebimento manual",
        payment_method=method,
        date=target_date,
        automatic=False
    ).first()

    if amount == 0:
        if existing:
            db.session.delete(existing)
    elif existing:
        existing.amount = amount
        existing.description = f"Recebimento manual - {method}"
    else:
        db.session.add(Transaction(
            kind="income",
            category="Recebimento manual",
            description=f"Recebimento manual - {method}",
            amount=amount,
            payment_method=method,
            date=target_date,
            automatic=False,
        ))

    db.session.commit()
    flash(
        f"Entrada manual de {method} atualizada para {brl(amount)}. Os totais foram recalculados.",
        "success"
    )
    return redirect(request.referrer or url_for("admin_cash", date=target_date.isoformat()))


@app.route("/admin/caixa/fechar", methods=["POST"])
@admin_required
def admin_close_cash():
    try:
        selected_date = datetime.strptime(request.form.get("date", ""), "%Y-%m-%d").date()
    except ValueError:
        flash("Data inválida.", "danger")
        return redirect(url_for("admin_cash"))

    cash = CashDay.query.filter_by(date=selected_date).first()
    if not cash:
        cash = CashDay(date=selected_date, opening_balance=0)
        db.session.add(cash)

    cash.closing_balance = decimal_from_form(request.form.get("closing_balance", "0"))
    cash.notes = request.form.get("notes", cash.notes or "").strip()
    cash.closed = True
    db.session.commit()

    flash("Caixa fechado.", "success")
    return redirect(url_for("admin_cash", date=selected_date.isoformat()))



@app.route("/admin/clientes/novo", methods=["POST"])
@admin_required
def admin_create_client():
    name = request.form.get("name", "").strip()
    phone = request.form.get("phone", "").strip()
    email = request.form.get("email", "").strip().lower()

    if not name:
        flash("Informe o nome do cliente.", "danger")
        return redirect(url_for("admin_clients"))

    if email and User.query.filter_by(email=email).first():
        flash("Já existe um cliente com esse e-mail.", "danger")
        return redirect(url_for("admin_clients"))

    if not email:
        email = f"cliente-{secrets.token_hex(6)}@cadastro.local"

    client = User(
        name=name,
        email=email,
        phone=phone,
        role="client",
        password_hash=""
    )
    client.set_password(secrets.token_urlsafe(24))
    db.session.add(client)
    db.session.flush()

    frequency = request.form.get("frequency", "").strip()
    if frequency in {"semanal", "quinzenal", "mensal", "bimestral"}:
        service_id = int(request.form.get("service_id", "0") or 0) or None
        plan_name = request.form.get("plan_name", "").strip() or "Plano recorrente"
        next_date_raw = request.form.get("next_date", "").strip()
        next_date = None
        if next_date_raw:
            try:
                next_date = datetime.strptime(next_date_raw, "%Y-%m-%d").date()
            except ValueError:
                next_date = None

        db.session.add(ClientPlan(
            user_id=client.id,
            service_id=service_id,
            plan_name=plan_name,
            frequency=frequency,
            amount=decimal_from_form(request.form.get("plan_amount", "0")),
            next_date=next_date,
            notes=request.form.get("plan_notes", "").strip(),
            active=True,
        ))

    db.session.commit()
    flash("Cliente cadastrado na administração.", "success")
    return redirect(url_for("admin_clients"))


@app.route("/admin/clientes/<int:client_id>/plano", methods=["POST"])
@admin_required
def admin_add_client_plan(client_id):
    client = db.session.get(User, client_id)
    if not client or client.role != "client":
        abort(404)

    frequency = request.form.get("frequency", "").strip()
    if frequency not in {"semanal", "quinzenal", "mensal", "bimestral"}:
        flash("Escolha uma frequência válida.", "danger")
        return redirect(url_for("admin_clients"))

    service_id = int(request.form.get("service_id", "0") or 0) or None
    next_date = None
    if request.form.get("next_date"):
        try:
            next_date = datetime.strptime(request.form.get("next_date"), "%Y-%m-%d").date()
        except ValueError:
            pass

    db.session.add(ClientPlan(
        user_id=client.id,
        service_id=service_id,
        plan_name=request.form.get("plan_name", "").strip() or "Plano recorrente",
        frequency=frequency,
        amount=decimal_from_form(request.form.get("amount", "0")),
        next_date=next_date,
        notes=request.form.get("notes", "").strip(),
        active=True,
    ))
    db.session.commit()
    flash("Plano recorrente adicionado.", "success")
    return redirect(url_for("admin_clients"))


@app.route("/admin/planos/<int:plan_id>/avancar", methods=["POST"])
@admin_required
def admin_advance_plan(plan_id):
    plan = db.session.get(ClientPlan, plan_id)
    if not plan:
        abort(404)

    plan.next_date = advance_recurring_date(plan.next_date or date.today(), plan.frequency)
    db.session.commit()
    flash("Próxima data do cliente atualizada.", "success")
    return redirect(url_for("admin_clients"))


@app.route("/admin/planos/<int:plan_id>/toggle", methods=["POST"])
@admin_required
def admin_toggle_plan(plan_id):
    plan = db.session.get(ClientPlan, plan_id)
    if not plan:
        abort(404)

    plan.active = not plan.active
    db.session.commit()
    return redirect(url_for("admin_clients"))


@app.route("/admin/clientes")
@admin_required
def admin_clients():
    q = request.args.get("q", "").strip()
    clients_query = User.query.filter_by(role="client")

    if q:
        like = f"%{q}%"
        clients_query = clients_query.filter(
            db.or_(
                User.name.ilike(like),
                User.email.ilike(like),
                User.phone.ilike(like),
            )
        )

    clients = clients_query.order_by(User.created_at.desc()).all()

    rows = []
    for client in clients:
        completed = Appointment.query.filter_by(
            user_id=client.id, status="concluido"
        ).all()
        total_spent = sum((Decimal(a.amount or 0) for a in completed), Decimal("0"))
        rows.append({
            "client": client,
            "visits": len(completed),
            "spent": total_spent,
            "vehicles": Vehicle.query.filter_by(user_id=client.id).all(),
        })

    services = Service.query.filter_by(active=True).order_by(Service.name.asc()).all()
    return render_template("admin_clients.html", rows=rows, q=q, services=services)


@app.route("/admin/funcionarios", methods=["GET", "POST"])
@admin_required
def admin_employees():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        if not name:
            flash("Informe o nome do funcionário.", "danger")
        else:
            emp = Employee(
                name=name,
                role_name=request.form.get("role_name", "Lavador").strip() or "Lavador",
                phone=request.form.get("phone", "").strip(),
                commission_percent=decimal_from_form(request.form.get("commission_percent", "0")),
                notes=request.form.get("notes", "").strip(),
                active=True,
            )
            db.session.add(emp)
            db.session.commit()
            flash("Funcionário cadastrado.", "success")
            return redirect(url_for("admin_employees"))

    employees = Employee.query.order_by(Employee.active.desc(), Employee.name.asc()).all()
    month_start, month_end, _ = month_bounds()

    employee_rows = []
    for emp in employees:
        completed = Appointment.query.filter(
            Appointment.employee_id == emp.id,
            Appointment.status == "concluido",
            Appointment.scheduled_date >= month_start,
            Appointment.scheduled_date <= month_end,
        ).all()
        generated = sum((Decimal(a.amount or 0) for a in completed), Decimal("0"))
        commission = (generated * Decimal(emp.commission_percent or 0) / Decimal("100")).quantize(Decimal("0.01"))
        employee_rows.append({
            "employee": emp,
            "services": len(completed),
            "generated": generated,
            "commission": commission,
        })

    return render_template("admin_employees.html", rows=employee_rows)


@app.route("/admin/funcionarios/<int:employee_id>/toggle", methods=["POST"])
@admin_required
def admin_toggle_employee(employee_id):
    emp = db.session.get(Employee, employee_id)
    if not emp:
        abort(404)
    emp.active = not emp.active
    db.session.commit()
    return redirect(url_for("admin_employees"))


@app.route("/admin/servicos", methods=["GET", "POST"])
@admin_required
def admin_services():
    if request.method == "POST":
        try:
            name = request.form.get("name", "").strip()
            price = decimal_from_form(request.form.get("price", "0"))
            duration = int(request.form.get("duration_minutes", "60"))

            if not name or price < 0:
                raise ValueError

            db.session.add(Service(
                name=name,
                description=request.form.get("description", "").strip(),
                price=price,
                duration_minutes=duration,
                active=True,
            ))
            db.session.commit()
            flash("Serviço adicionado.", "success")
            return redirect(url_for("admin_services"))
        except Exception:
            db.session.rollback()
            flash("Confira os dados do serviço.", "danger")

    services = Service.query.order_by(Service.name.asc()).all()
    return render_template("admin_services.html", services=services)


@app.route("/admin/servicos/<int:service_id>/preco", methods=["POST"])
@admin_required
def admin_edit_service_price(service_id):
    service = db.session.get(Service, service_id)
    if not service:
        abort(404)

    price = decimal_from_form(request.form.get("price", "0"))
    if price < 0:
        flash("O valor do serviço não pode ser negativo.", "danger")
        return redirect(url_for("admin_services"))

    service.price = price
    db.session.commit()
    flash(f"Valor de {service.name} atualizado para {brl(price)}.", "success")
    return redirect(url_for("admin_services"))


@app.route("/admin/servicos/<int:service_id>/toggle", methods=["POST"])
@admin_required
def admin_toggle_service(service_id):
    service = db.session.get(Service, service_id)
    if not service:
        abort(404)
    service.active = not service.active
    db.session.commit()
    return redirect(url_for("admin_services"))


@app.route("/admin/relatorios")
@admin_required
def admin_reports():
    start_date, end_date, selected_month = month_bounds(request.args.get("month", ""))

    income, expenses, profit = transaction_totals(start_date, end_date)

    completed = Appointment.query.filter(
        Appointment.status == "concluido",
        Appointment.scheduled_date >= start_date,
        Appointment.scheduled_date <= end_date,
    ).all()

    ticket_avg = (sum((Decimal(a.amount or 0) for a in completed), Decimal("0")) / len(completed)) if completed else Decimal("0")

    service_summary = db.session.query(
        Service.name,
        db.func.count(Appointment.id),
        db.func.coalesce(db.func.sum(Appointment.amount), 0)
    ).join(Appointment, Appointment.service_id == Service.id).filter(
        Appointment.status == "concluido",
        Appointment.scheduled_date >= start_date,
        Appointment.scheduled_date <= end_date,
    ).group_by(Service.id, Service.name).order_by(db.func.sum(Appointment.amount).desc()).all()

    expense_summary = db.session.query(
        Transaction.category,
        db.func.coalesce(db.func.sum(Transaction.amount), 0)
    ).filter(
        Transaction.kind == "expense",
        Transaction.date >= start_date,
        Transaction.date <= end_date,
    ).group_by(Transaction.category).order_by(db.func.sum(Transaction.amount).desc()).all()

    commission_summary = db.session.query(
        Employee.name,
        db.func.count(Appointment.id),
        db.func.coalesce(db.func.sum(Appointment.amount), 0),
        Employee.commission_percent
    ).join(Appointment, Appointment.employee_id == Employee.id).filter(
        Appointment.status == "concluido",
        Appointment.scheduled_date >= start_date,
        Appointment.scheduled_date <= end_date,
    ).group_by(Employee.id, Employee.name, Employee.commission_percent).all()

    commissions = []
    for name, count, generated, pct in commission_summary:
        generated = Decimal(generated or 0)
        pct = Decimal(pct or 0)
        commissions.append((name, count, generated, (generated * pct / Decimal("100")).quantize(Decimal("0.01"))))

    return render_template(
        "admin_reports.html",
        selected_month=selected_month,
        start_date=start_date,
        end_date=end_date,
        income=income,
        expenses=expenses,
        profit=profit,
        completed_count=len(completed),
        ticket_avg=ticket_avg,
        service_summary=service_summary,
        expense_summary=expense_summary,
        commissions=commissions,
    )


@app.route("/admin/relatorios/exportar.xlsx")
@admin_required
def admin_reports_xlsx():
    start_date, end_date, selected_month = month_bounds(request.args.get("month", ""))

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"

    income, expenses, profit = transaction_totals(start_date, end_date)
    completed = Appointment.query.filter(
        Appointment.status == "concluido",
        Appointment.scheduled_date >= start_date,
        Appointment.scheduled_date <= end_date,
    ).all()

    ws.append(["BOSTON CAR WASH - RELATÓRIO MENSAL"])
    ws.append(["Período", f"{start_date.strftime('%d/%m/%Y')} a {end_date.strftime('%d/%m/%Y')}"])
    ws.append([])
    ws.append(["Receitas", float(income)])
    ws.append(["Despesas", float(expenses)])
    ws.append(["Lucro", float(profit)])
    ws.append(["Serviços concluídos", len(completed)])
    ws.append(["Ticket médio", float(sum((Decimal(a.amount or 0) for a in completed), Decimal('0')) / len(completed)) if completed else 0])

    ws["A1"].font = Font(size=16, bold=True)
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 25

    finance = wb.create_sheet("Financeiro")
    finance.append(["Data", "Tipo", "Categoria", "Descrição", "Forma de pagamento", "Valor", "Automático"])
    txs = Transaction.query.filter(
        Transaction.date >= start_date,
        Transaction.date <= end_date
    ).order_by(Transaction.date.asc(), Transaction.id.asc()).all()

    for tx in txs:
        finance.append([
            tx.date.strftime("%d/%m/%Y"),
            "Receita" if tx.kind == "income" else "Despesa",
            tx.category,
            tx.description,
            tx.payment_method,
            float(tx.amount),
            "Sim" if tx.automatic else "Não",
        ])

    appointments = wb.create_sheet("Atendimentos")
    appointments.append(["Data", "Hora", "Cliente", "Veículo", "Placa", "Serviço", "Funcionário", "Status", "Valor"])
    for a in Appointment.query.filter(
        Appointment.scheduled_date >= start_date,
        Appointment.scheduled_date <= end_date
    ).order_by(Appointment.scheduled_date.asc(), Appointment.scheduled_time.asc()).all():
        appointments.append([
            a.scheduled_date.strftime("%d/%m/%Y"),
            a.scheduled_time,
            a.client.name,
            a.vehicle,
            a.plate or "",
            a.service.name,
            a.employee.name if a.employee else "",
            a.status.replace("_", " "),
            float(a.amount),
        ])

    commissions = wb.create_sheet("Comissões")
    commissions.append(["Funcionário", "Serviços concluídos", "Valor gerado", "% Comissão", "Comissão"])
    for emp in Employee.query.order_by(Employee.name.asc()).all():
        emp_appts = Appointment.query.filter(
            Appointment.employee_id == emp.id,
            Appointment.status == "concluido",
            Appointment.scheduled_date >= start_date,
            Appointment.scheduled_date <= end_date,
        ).all()
        generated = sum((Decimal(a.amount or 0) for a in emp_appts), Decimal("0"))
        commission = generated * Decimal(emp.commission_percent or 0) / Decimal("100")
        commissions.append([emp.name, len(emp_appts), float(generated), float(emp.commission_percent or 0), float(commission)])

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0B5E8E")
            cell.alignment = Alignment(horizontal="center")
        for col in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in col) + 2, 45)
            sheet.column_dimensions[col[0].column_letter].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=f"Boston_Car_Wash_Relatorio_{selected_month}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def seed():
    db.create_all()

    # Migração simples para quem reaproveitar um boston.db de uma versão anterior.
    if app.config["SQLALCHEMY_DATABASE_URI"].startswith("sqlite"):
        with db.engine.begin() as conn:
            columns = {
                row[1] for row in conn.execute(text("PRAGMA table_info(appointment)")).fetchall()
            }
            if "payment_method" not in columns:
                conn.execute(text(
                    "ALTER TABLE appointment ADD COLUMN payment_method VARCHAR(50) "
                    "NOT NULL DEFAULT 'Não informado'"
                ))
            if "completed_at" not in columns:
                conn.execute(text(
                    "ALTER TABLE appointment ADD COLUMN completed_at DATETIME"
                ))

    admin_email = os.getenv("ADMIN_EMAIL", "admin@bostoncarwash.local").strip().lower()
    admin_password = os.getenv("ADMIN_PASSWORD", "TroqueEssaSenha123!")

    if not User.query.filter_by(email=admin_email).first():
        admin = User(
            name="Administrador",
            email=admin_email,
            role="admin",
            password_hash=""
        )
        admin.set_password(admin_password)
        db.session.add(admin)

    if Service.query.count() == 0:
        defaults = [
            ("Lavagem Completa", "Limpeza externa completa com acabamento caprichado.", 80, 60),
            ("Ducha com Cera", "Lavagem com finalização em cera para brilho e proteção.", 100, 75),
            ("Higienização Interna", "Limpeza profunda do interior do veículo.", 250, 180),
            ("Detalhamento de Rodas", "Limpeza e acabamento detalhado de rodas.", 120, 90),
            ("Limpeza Técnica do Motor", "Limpeza técnica cuidadosa da região do motor.", 180, 120),
            ("Revitalização de Plásticos", "Tratamento para recuperar aparência de plásticos.", 150, 90),
        ]
        for name, desc, price, duration in defaults:
            db.session.add(Service(
                name=name,
                description=desc,
                price=price,
                duration_minutes=duration,
                active=True,
            ))

    db.session.commit()


with app.app_context():
    seed()


if __name__ == "__main__":
    app.run(debug=True)
